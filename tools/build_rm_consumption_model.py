"""Build RM_Consumption_Model_2026.xlsx from nine monthly inventory-movement exports.
All Model/Summary cells are live Excel formulas referencing the M01..M09 raw sheets."""
import sys, datetime
import openpyxl
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import os
HERE = os.path.dirname(os.path.abspath(__file__))
# usage: python3 tools/build_rm_consumption_model.py [input_dir] [output_xlsx]
IN_DIR = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "data", "rm")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "..", "output", "RM_Consumption_Model_2026.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
NM = 9
MONTH_NAMES = ["كانون الثاني Jan", "شباط Feb", "آذار Mar", "نيسان Apr", "أيار May", "حزيران Jun",
               "تموز Jul", "آب Aug", "أيلول Sep"]
# source column letters in the month sheets
S = dict(ItemNo="C", ItemDesc="D", SubCateg="A", OpenQty="E", OpenCost="F", TotalInQty="M",
         TotalOutQty="S", EndQty="T", EndCost="U", ItemUnitCost="V", ConsMonRate9="AD", PurchAmt="AG")
DATE_COLS = {27, 29}  # Last Purchase Date, LastIssueSaleDate

wb = openpyxl.Workbook()
wb.remove(wb.active)

hdr_font = Font(bold=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---------- 1. raw month sheets ----------
nrows = None
first_item_rows = None
for m in range(1, NM + 1):
    src = openpyxl.load_workbook(f"{IN_DIR}/rm_{m}.xlsx", data_only=True)["Sheet"]
    ws = wb.create_sheet(f"M{m:02d}")
    for r, row in enumerate(src.iter_rows(values_only=True), start=1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if r == 1:
                cell.font = hdr_font
            elif c in DATE_COLS and isinstance(v, datetime.datetime):
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in range(1, 34):
        ws.column_dimensions[L(c)].width = 14
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["G"].width = 18
    if nrows is None:
        nrows = src.max_row
        first_item_rows = [src.cell(row=r, column=3).value for r in range(2, nrows + 1)]
    else:
        assert src.max_row == nrows
        assert [src.cell(row=r, column=3).value for r in range(2, nrows + 1)] == first_item_rows, f"ItemNo order differs in M{m:02d}"
# source files carry a totals row at the bottom (ItemNo blank): kept verbatim in M01..M09, excluded from Model
LAST = max(i + 2 for i, v in enumerate(first_item_rows) if v is not None)  # last item row (header is row 1)
print("item rows: 2 ..", LAST, "(sheet rows incl. totals row:", nrows, ")")

# ---------- 2. Model ----------
ws = wb.create_sheet("Model", 0)
FILL_A = PatternFill("solid", fgColor="DDEBF7")
FILL_B = PatternFill("solid", fgColor="E2EFDA")
FILL_S = PatternFill("solid", fgColor="FFF2CC")
FILL_K = PatternFill("solid", fgColor="F8CBAD")
FILL_ID = PatternFill("solid", fgColor="EDEDED")

BLOCK = 12
FIRST_BLOCK_COL = 4
def bcol(m, off):  # column index of offset `off` (0..11) within month m block
    return FIRST_BLOCK_COL + (m - 1) * BLOCK + off
SUM_COL = bcol(NM + 1, 0)  # first summary column

month_hdrs = [
    "رصيد بداية المدة كمية\nOpenQty", "كلفة بداية المدة\nOpenCost", "كلفة الوحدة بداية\nOpen Unit Cost",
    "مجموع الإدخال كمية\nTotalInQty", "كلفة الإدخال\nInCost", "كلفة وحدة الإدخال\nIn Unit Cost",
    "مجموع الإخراج كمية\nOutQty (+)", "كلفة الإخراج\nOutCost", "كلفة وحدة الإخراج\nOut Unit Cost",
    "رصيد نهاية المدة كمية\nEndQty", "كلفة نهاية المدة\nEndCost", "كلفة الوحدة نهاية\nEnd Unit Cost",
]
sum_hdrs = [
    "إجمالي الإخراج 9 أشهر كمية\nTotal Out Qty 9M",
    "إجمالي كلفة الإخراج 9 أشهر\nTotal Out Cost 9M",
    "متوسط الاستهلاك الشهري الفعلي كمية\nAvg Monthly Out Qty",
    "متوسط كلفة الاستهلاك الشهري\nAvg Monthly Out Cost",
    "متوسط كلفة الوحدة الفعلية 9 أشهر\nActual Avg Unit Cost 9M",
    "كلفة الوحدة المعيارية في النظام\nStd Unit Cost (ItemUnitCost M09)",
    "الفرق % بين الفعلية والمعيارية\nActual vs Std %",
    "معدل الاستهلاك الشهري في النظام\nSystem ConsMonRate9 (M09)",
    "الفرق بين معدلنا ومعدل النظام\nOur Rate - System Rate",
    "عدد الأشهر التي فيها حركة إخراج\nMonths with Out Movement",
    "رصيد نهاية أيلول كمية\nEndQty Sep",
    "شهور التغطية\nCoverage Months",
    "تدقيق (يجب أن يكون 0)\nCheck (=0)",
]
# header rows
ws.cell(row=1, column=1, value="بيانات المادة / Item")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
for c, h in enumerate(["رقم المادة\nItemNo", "وصف المادة\nItemDesc", "الفئة الفرعية\nSubCateg"], start=1):
    ws.cell(row=2, column=c, value=h).fill = FILL_ID
    ws.cell(row=1, column=c).fill = FILL_ID
for m in range(1, NM + 1):
    fill = FILL_A if m % 2 else FILL_B
    c0 = bcol(m, 0)
    ws.cell(row=1, column=c0, value=f"M{m:02d} - {MONTH_NAMES[m-1]} 2026")
    ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + BLOCK - 1)
    for off, h in enumerate(month_hdrs):
        ws.cell(row=2, column=c0 + off, value=h).fill = fill
        ws.cell(row=1, column=c0 + off).fill = fill
ws.cell(row=1, column=SUM_COL, value="الملخص 9 أشهر (كانون الثاني - أيلول 2026) / 9-Month Summary")
ws.merge_cells(start_row=1, start_column=SUM_COL, end_row=1, end_column=SUM_COL + len(sum_hdrs) - 1)
for off, h in enumerate(sum_hdrs):
    f = FILL_K if off == len(sum_hdrs) - 1 else FILL_S
    ws.cell(row=2, column=SUM_COL + off, value=h).fill = f
    ws.cell(row=1, column=SUM_COL + off).fill = f
for c in range(1, SUM_COL + len(sum_hdrs)):
    for r in (1, 2):
        cell = ws.cell(row=r, column=c)
        cell.font = hdr_font
        cell.alignment = center
        cell.border = border
ws.row_dimensions[2].height = 60

AMT = "#,##0.00"
UNIT = "0.0000"
PCT = "0.0%"
fmt_month = [AMT, AMT, UNIT, AMT, AMT, UNIT, AMT, AMT, UNIT, AMT, AMT, UNIT]
fmt_sum = [AMT, AMT, AMT, AMT, UNIT, UNIT, PCT, AMT, AMT, "0", AMT, "0.00", AMT]

representative = {}
MODEL_LAST = LAST + 1  # Model has two header rows: source row r -> Model row r+1
for r in range(2, LAST + 1):
    R = r + 1  # model row
    ws.cell(row=R, column=1, value=f"='M01'!{S['ItemNo']}{r}")
    ws.cell(row=R, column=2, value=f"='M01'!{S['ItemDesc']}{r}")
    ws.cell(row=R, column=3, value=f"='M01'!{S['SubCateg']}{r}")
    for m in range(1, NM + 1):
        sh = f"'M{m:02d}'"
        c = [L(bcol(m, o)) for o in range(BLOCK)]  # letters of block columns
        oq, oc, ouc, iq, ic, iuc, xq, xc, xuc, eq, ec, euc = [f"{x}{R}" for x in c]
        f = [
            f"={sh}!{S['OpenQty']}{r}",
            f"={sh}!{S['OpenCost']}{r}",
            f"=IFERROR({oc}/{oq},0)",
            f"={sh}!{S['TotalInQty']}{r}",
            f"=IF({sh}!{S['PurchAmt']}{r}>0,{sh}!{S['PurchAmt']}{r},{iq}*IFERROR({oc}/{oq},{sh}!{S['ItemUnitCost']}{r}))",
            f"=IFERROR({ic}/{iq},0)",
            f"=-{sh}!{S['TotalOutQty']}{r}",
            f"={oc}+{ic}-{ec}",
            f"=IFERROR({xc}/{xq},0)",
            f"={sh}!{S['EndQty']}{r}",
            f"={sh}!{S['EndCost']}{r}",
            f"=IFERROR({ec}/{eq},0)",
        ]
        for o, fx in enumerate(f):
            cell = ws.cell(row=R, column=bcol(m, o), value=fx)
            cell.number_format = fmt_month[o]
        if r == 2 and m == 2:
            representative["month2"] = f
    # summary
    in_cells = ",".join(f"{L(bcol(m, 3))}{R}" for m in range(1, NM + 1))
    out_cells = ",".join(f"{L(bcol(m, 6))}{R}" for m in range(1, NM + 1))
    outcost_cells = ",".join(f"{L(bcol(m, 7))}{R}" for m in range(1, NM + 1))
    countif = "+".join(f'COUNTIF({L(bcol(m, 6))}{R},">0")' for m in range(1, NM + 1))
    sc = [f"{L(SUM_COL + o)}{R}" for o in range(len(sum_hdrs))]
    tq, tc, aq, ac, auc, std, pct, sysr, dr, nmv, eqs, cov, chk = sc
    end_sep = f"{L(bcol(NM, 9))}{R}"
    open_jan = f"{L(bcol(1, 0))}{R}"
    fs = [
        f"=SUM({out_cells})",
        f"=SUM({outcost_cells})",
        f"={tq}/{NM}",
        f"={tc}/{NM}",
        f"=IFERROR({tc}/{tq},0)",
        f"='M{NM:02d}'!{S['ItemUnitCost']}{r}",
        f"=IFERROR({auc}/{std}-1,0)",
        f"='M{NM:02d}'!{S['ConsMonRate9']}{r}",
        f"={aq}-{sysr}",
        f"={countif}",
        f"={end_sep}",
        f"=IFERROR({eqs}/{aq},0)",
        f"={open_jan}+SUM({in_cells})-SUM({out_cells})-{end_sep}",
    ]
    for o, fx in enumerate(fs):
        cell = ws.cell(row=R, column=SUM_COL + o, value=fx)
        cell.number_format = fmt_sum[o]
    if r == 2:
        representative["summary"] = fs

ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A2:{L(SUM_COL + len(sum_hdrs) - 1)}{MODEL_LAST}"
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 32
ws.column_dimensions["C"].width = 10
for c in range(FIRST_BLOCK_COL, SUM_COL + len(sum_hdrs)):
    ws.column_dimensions[L(c)].width = 14
for o in range(len(sum_hdrs)):
    ws.column_dimensions[L(SUM_COL + o)].width = 16

# ---------- 3. Summary per SubCateg ----------
cats = sorted({v for v in (wb["M01"].cell(row=r, column=1).value for r in range(2, LAST + 1)) if v is not None})
ss = wb.create_sheet("Summary", 1)
sh_ = ["الفئة الفرعية\nSubCateg", "عدد المواد\nItems", "عدد المواد التي لها إخراج\nItems with Out Movement",
       "إجمالي كمية الإخراج 9 أشهر\nTotal Out Qty 9M", "إجمالي كلفة الإخراج 9 أشهر\nTotal Out Cost 9M",
       "متوسط كلفة الاستهلاك الشهري\nAvg Monthly Out Cost", "حصة من الإجمالي %\nShare of Total %",
       "مجموع التدقيق (يجب أن يكون 0)\nSum of Check"]
for c, h in enumerate(sh_, start=1):
    cell = ss.cell(row=1, column=c, value=h)
    cell.font = hdr_font; cell.alignment = center; cell.fill = FILL_S; cell.border = border
ss.row_dimensions[1].height = 48
rngC = f"Model!$C$3:$C${MODEL_LAST}"
colTQ, colTC, colAC, colNMV, colCHK = (f"Model!${L(SUM_COL+o)}$3:${L(SUM_COL+o)}${MODEL_LAST}" for o in (0, 1, 3, 9, 12))
n = len(cats)
tot_row = n + 2
for i, cat in enumerate(cats):
    r = i + 2
    ss.cell(row=r, column=1, value=cat)
    ss.cell(row=r, column=2, value=f'=COUNTIF({rngC},A{r})').number_format = "0"
    ss.cell(row=r, column=3, value=f'=COUNTIFS({rngC},A{r},{colNMV},">0")').number_format = "0"
    ss.cell(row=r, column=4, value=f"=SUMIF({rngC},A{r},{colTQ})").number_format = AMT
    ss.cell(row=r, column=5, value=f"=SUMIF({rngC},A{r},{colTC})").number_format = AMT
    ss.cell(row=r, column=6, value=f"=SUMIF({rngC},A{r},{colAC})").number_format = AMT
    ss.cell(row=r, column=7, value=f"=IFERROR(E{r}/$E${tot_row},0)").number_format = PCT
    ss.cell(row=r, column=8, value=f"=SUMIF({rngC},A{r},{colCHK})").number_format = AMT
ss.cell(row=tot_row, column=1, value="الإجمالي العام / Grand Total")
for c, fx in zip(range(2, 9), [f"=SUM(B2:B{n+1})", f"=SUM(C2:C{n+1})", f"=SUM(D2:D{n+1})", f"=SUM(E2:E{n+1})",
                               f"=SUM(F2:F{n+1})", f"=SUM(G2:G{n+1})", f"=SUM(H2:H{n+1})"]):
    ss.cell(row=tot_row, column=c, value=fx).number_format = PCT if c == 7 else (AMT if c > 3 else "0")
for c in range(1, 9):
    ss.cell(row=tot_row, column=c).font = hdr_font
    ss.cell(row=tot_row, column=c).fill = FILL_ID
    ss.column_dimensions[L(c)].width = 20
ss.freeze_panes = "B2"

# ---------- 4. README ----------
rd = wb.create_sheet("README", 2)
notes = [
    ("نموذج الاستهلاك الشهري الفعلي للمواد الخام 2026", "Raw-Material Actual Monthly Consumption Model 2026"),
    ("", ""),
    ("مصدر البيانات: تسعة ملفات حركة مخزون شهرية rm_1.xlsx .. rm_9.xlsx (562 مادة في كل ملف، نفس الترتيب).",
     "Source: nine monthly inventory-movement exports rm_1.xlsx .. rm_9.xlsx (562 items each, identical order)."),
    ("الأوراق M01..M09 هي نسخة حرفية من البيانات الأصلية (33 عموداً) ولا يجب تعديلها؛ كل خلايا ورقة Model وSummary صيغ تشير إليها.",
     "Sheets M01..M09 are verbatim copies of the raw data (33 columns); do not edit them. Every cell in Model and Summary is a live formula referencing them."),
    ("تطابق الأشهر: M01 = كانون الثاني 2026، M02 = شباط، M03 = آذار، M04 = نيسان، M05 = أيار، M06 = حزيران، M07 = تموز، M08 = آب، M09 = أيلول 2026.",
     "Month mapping: M01 = Jan 2026, M02 = Feb, M03 = Mar, M04 = Apr, M05 = May, M06 = Jun, M07 = Jul, M08 = Aug, M09 = Sep 2026."),
    ("", ""),
    ("قواعد الكلفة:", "Cost rules:"),
    ("كلفة وحدة البداية = IFERROR(OpenCost / OpenQty, 0)", "Open unit cost = IFERROR(OpenCost / OpenQty, 0)"),
    ("كلفة الإدخال = IF(Purchase Amount > 0, Purchase Amount, TotalInQty × IFERROR(OpenCost/OpenQty, ItemUnitCost)) — قيمة الشراء الفعلية، وعند عدم وجود شراء تُقدَّر بكلفة الوحدة الافتتاحية (أو المعيارية إذا كان الرصيد الافتتاحي صفراً).",
     "InCost = IF(Purchase Amount > 0, Purchase Amount, TotalInQty × IFERROR(OpenCost/OpenQty, ItemUnitCost)) — actual purchase value; non-purchase inflows are valued at the opening unit cost (or the standard cost when opening stock is zero)."),
    ("كلفة الإخراج = OpenCost + InCost − EndCost (رقم موازنة = الكلفة الفعلية بالمتوسط المتحرك لما خرج من المخزون).",
     "OutCost = OpenCost + InCost − EndCost (balancing figure = actual moving-average cost of what left stock)."),
    ("كلفة وحدة الإخراج = IFERROR(OutCost / OutQty, 0)؛ كلفة وحدة النهاية = IFERROR(EndCost / EndQty, 0).",
     "Out unit cost = IFERROR(OutCost / OutQty, 0); End unit cost = IFERROR(EndCost / EndQty, 0)."),
    ("", ""),
    ("ملاحظات:", "Notes:"),
    ("TotalOutQty سالب في ملفات المصدر؛ يُعرض في النموذج بقيمة موجبة (= −TotalOutQty).",
     "TotalOutQty is negative in the source files; the model shows it as a positive quantity (= −TotalOutQty)."),
    ("ItemUnitCost في النظام كلفة معيارية ثابتة عبر الأشهر وليست الكلفة الفعلية؛ الكلفة الفعلية للوحدة = الكلفة ÷ الكمية.",
     "ItemUnitCost in the system is a constant standard cost across months, NOT actual; actual unit cost = Cost / Qty."),
    ("الخلايا الفارغة في المصدر تُعامل كصفر في الصيغ.", "Blank numeric cells in the source are treated as 0 by the formulas."),
    ("عمود التدقيق Check = OpenQty(M01) + مجموع الإدخال 9 أشهر − مجموع الإخراج 9 أشهر − EndQty(M09) ويجب أن يساوي صفراً لكل مادة.",
     "Check column = OpenQty(M01) + SUM(9 in qty) − SUM(9 out qty) − EndQty(M09); must be 0 for every item."),
    ("شهور التغطية = رصيد نهاية أيلول ÷ متوسط الإخراج الشهري الفعلي.", "Coverage months = EndQty Sep / actual average monthly out qty."),
    ("ورقة Summary: إجماليات لكل فئة فرعية SubCateg بصيغ SUMIF/COUNTIF على ورقة Model.",
     "Summary sheet: per-SubCateg totals via SUMIF/COUNTIF formulas on the Model sheet."),
    ("الملف لا يحتوي قيماً محسوبة مخزنة؛ يحسب Excel كل الصيغ عند الفتح.",
     "The file stores no cached values; Excel calculates all formulas on open."),
]
rd.cell(row=1, column=1, value="العربية").font = hdr_font
rd.cell(row=1, column=2, value="English").font = hdr_font
for i, (a, e) in enumerate(notes, start=2):
    ca = rd.cell(row=i, column=1, value=a); ce = rd.cell(row=i, column=2, value=e)
    ca.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")
    ce.alignment = Alignment(wrap_text=True, vertical="top")
    if i == 2:
        ca.font = ce.font = Font(bold=True, size=13)
rd.column_dimensions["A"].width = 90
rd.column_dimensions["B"].width = 90

wb.save(OUT)
print("saved", OUT)
print("rows:", LAST, "summary first col:", L(SUM_COL), "last col:", L(SUM_COL + len(sum_hdrs) - 1))
print("block ranges:", [(f"M{m:02d}", L(bcol(m,0)), L(bcol(m,11))) for m in range(1, NM+1)])
print("REPRESENTATIVE month2 row2:"); [print("  ", x) for x in representative["month2"]]
print("REPRESENTATIVE summary row2:"); [print("  ", x) for x in representative["summary"]]
