from datetime import datetime, date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _num(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _has_item(row: tuple) -> bool:
    item_no = row[0]
    if isinstance(item_no, str) and item_no.strip():
        return True
    return False


def parse_invoice(path: Path) -> dict:
    """Parse a BakeMark-style invoice workbook.

    Sheet "الفواتير" layout:
      Row 1: OrderNo (col B), OrderDate (col H)
      Row 2: Customer Name (col B)
      Row 4: headers (ItemNo, ItemDesc_Ara, Qty, U.Price, U.Cost, O.H,
                      Total Cost, Total Value, Margin, Margin %)
      Row 5+: pairs of rows per item — first row in USD, second in IQD
      Last data row before EOF: "TOTAL INV" with grand totals (USD then IQD).
    """
    wb = load_workbook(path, data_only=True)
    ws = wb["الفواتير"]
    rows = list(ws.iter_rows(values_only=True))

    order_no = rows[0][1]
    order_date = rows[0][7]
    if isinstance(order_date, datetime):
        order_date = order_date.date()
    customer = rows[1][1]

    lines: list[dict] = []
    total: dict | None = None

    i = 4
    while i < len(rows):
        row = rows[i]

        if row and row[5] == "TOTAL INV":
            iqd_row = rows[i + 1] if i + 1 < len(rows) else (None,) * 10
            total = {
                "usd": {
                    "total_cost": _num(row[6]),
                    "total_value": _num(row[7]),
                    "margin": _num(row[8]),
                    "margin_pct": _num(row[9]),
                },
                "iqd": {
                    "total_cost": _num(iqd_row[6]),
                    "total_value": _num(iqd_row[7]),
                    "margin": _num(iqd_row[8]),
                },
            }
            break

        if _has_item(row):
            iqd_row = rows[i + 1] if i + 1 < len(rows) else (None,) * 10
            lines.append(
                {
                    "item_no": row[0],
                    "desc": row[1],
                    "qty": _num(row[2]),
                    "usd": {
                        "unit_price": _num(row[3]),
                        "unit_cost": _num(row[4]),
                        "total_cost": _num(row[6]),
                        "total_value": _num(row[7]),
                        "margin": _num(row[8]),
                        "margin_pct": _num(row[9]),
                    },
                    "iqd": {
                        "unit_price": _num(iqd_row[3]),
                        "unit_cost": _num(iqd_row[4]),
                        "total_cost": _num(iqd_row[6]),
                        "total_value": _num(iqd_row[7]),
                        "margin": _num(iqd_row[8]),
                    },
                }
            )
            i += 2
            continue

        i += 1

    if total is None:
        total = {
            "usd": {
                "total_cost": sum(it["usd"]["total_cost"] for it in lines),
                "total_value": sum(it["usd"]["total_value"] for it in lines),
                "margin": sum(it["usd"]["margin"] for it in lines),
                "margin_pct": 0.0,
            },
            "iqd": {
                "total_cost": sum(it["iqd"]["total_cost"] for it in lines),
                "total_value": sum(it["iqd"]["total_value"] for it in lines),
                "margin": sum(it["iqd"]["margin"] for it in lines),
            },
        }
        if total["usd"]["total_value"]:
            total["usd"]["margin_pct"] = total["usd"]["margin"] / total["usd"]["total_value"]

    return {
        "file": path.name,
        "order_no": order_no,
        "order_date": order_date,
        "customer": customer,
        "lines": lines,
        "total": total,
    }


def parse_all(data_dir: Path) -> list[dict]:
    invoices = []
    for path in sorted(data_dir.glob("*.xlsx")):
        try:
            invoices.append(parse_invoice(path))
        except Exception as exc:
            print(f"failed to parse {path.name}: {exc}")
    return invoices
